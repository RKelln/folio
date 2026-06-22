"""Metadata stripper for rendered corpus artifacts.

PURPOSE
    Authoring metadata (author, last-modified-by, PDF Producer/Creator, XMP
    toolkit) is a PII risk: it can leak the real name of whoever generated a
    file. After a corpus artifact is rendered, ``strip_metadata`` removes that
    metadata in place; ``read_metadata`` reads it back for tests/verification.

DISPATCH
    By file suffix:
      - .pdf / .png / .jpg / .jpeg -> exiftool ``-all=`` (REQUIRED; if exiftool
        is missing we RAISE rather than silently pass, because leaking PDF/image
        producer/author metadata defeats the purpose of the corpus).
      - .docx -> clear python-docx core properties + save, then exiftool if
        available.
      - .xlsx -> clear openpyxl workbook properties + save, then exiftool if
        available.
      - anything else -> no-op (returns False, logs a warning).
"""

from __future__ import annotations

import json
import logging
import shutil
import subprocess
from pathlib import Path

logger = logging.getLogger(__name__)

# Formats where stripping is mandatory and only exiftool can do it.
_EXIFTOOL_REQUIRED_SUFFIXES = {".pdf", ".png", ".jpg", ".jpeg"}

# Office authoring properties to blank, by format.
_DOCX_PROPS = (
    "author",
    "last_modified_by",
    "comments",
    "title",
    "subject",
    "keywords",
    "category",
    "content_status",
    "identifier",
)
_XLSX_PROPS = (
    "creator",
    "lastModifiedBy",
    "last_modified_by",
    "title",
    "subject",
    "keywords",
    "description",
    "category",
    "identifier",
)


def _exiftool_available() -> bool:
    """True when the ``exiftool`` binary is on PATH."""
    return shutil.which("exiftool") is not None


def _run_exiftool_strip(path: Path) -> None:
    """Run ``exiftool -all= -overwrite_original -q`` on ``path``.

    Raises:
        RuntimeError: If exiftool is missing or exits non-zero.
    """
    binary = shutil.which("exiftool")
    if binary is None:
        raise RuntimeError(
            f"exiftool not available to strip metadata from {path}; "
            f"refusing to leave authoring metadata in place."
        )
    try:
        subprocess.run(
            [binary, "-all=", "-overwrite_original", "-q", str(path)],
            capture_output=True,
            check=True,
        )
    except (subprocess.CalledProcessError, OSError) as exc:
        detail = getattr(exc, "stderr", b"") or b""
        raise RuntimeError(
            f"exiftool failed to strip {path}: "
            f"{detail.decode('utf-8', 'replace').strip()[:300] or exc}"
        ) from exc


def _try_exiftool_strip(path: Path) -> None:
    """Best-effort exiftool strip for office files (defense-in-depth).

    exiftool cannot *write* OOXML (.docx/.xlsx) containers, so a failure here is
    expected on some platforms and must NOT fail the strip: the office library
    has already cleared the authoritative core properties. The failure is logged
    (never silently swallowed) per AGENTS.md.
    """
    if not _exiftool_available():
        return
    try:
        _run_exiftool_strip(path)
    except RuntimeError as exc:
        logger.warning("best-effort exiftool strip skipped for %s: %s", path, exc)


def _strip_docx(path: Path) -> bool:
    """Clear python-docx core properties, save, then exiftool if available."""
    try:
        import docx  # type: ignore[import-not-found]
    except ImportError as exc:
        raise RuntimeError(
            "python-docx not installed; cannot strip .docx metadata."
        ) from exc
    document = docx.Document(str(path))
    cp = document.core_properties
    for attr in _DOCX_PROPS:
        try:
            setattr(cp, attr, "")
        except (ValueError, TypeError) as exc:
            logger.warning("Could not clear docx property %s on %s: %s", attr, path, exc)
    document.save(str(path))
    _try_exiftool_strip(path)
    return True


def _strip_xlsx(path: Path) -> bool:
    """Clear openpyxl workbook properties, save, then exiftool if available."""
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ImportError as exc:
        raise RuntimeError(
            "openpyxl not installed; cannot strip .xlsx metadata."
        ) from exc
    workbook = openpyxl.load_workbook(str(path))
    props = workbook.properties
    # Empty strings (not None) — openpyxl re-applies its default creator when a
    # property is None.
    for attr in _XLSX_PROPS:
        try:
            setattr(props, attr, "")
        except (ValueError, TypeError, AttributeError) as exc:
            logger.warning("Could not clear xlsx property %s on %s: %s", attr, path, exc)
    workbook.save(str(path))
    _try_exiftool_strip(path)
    return True


def strip_metadata(path: str | Path) -> bool:
    """Strip authoring metadata from a rendered artifact in place.

    Args:
        path: The file to strip. Dispatch is by suffix (see module docstring).

    Returns:
        True if metadata was stripped; False for unsupported suffixes.

    Raises:
        RuntimeError: For pdf/image formats when exiftool is unavailable, or if
            the office library required by docx/xlsx is missing, or if exiftool
            itself fails. We never silently skip a required strip.

    Side effects:
        Rewrites the file in place.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in _EXIFTOOL_REQUIRED_SUFFIXES:
        _run_exiftool_strip(path)
        return True
    if suffix == ".docx":
        return _strip_docx(path)
    if suffix == ".xlsx":
        return _strip_xlsx(path)
    logger.warning("strip_metadata: no metadata handler for %s; leaving %s unchanged", suffix, path)
    return False


def _read_exiftool(path: Path) -> dict:
    """Read metadata via ``exiftool -j``; returns {} when unavailable/empty."""
    binary = shutil.which("exiftool")
    if binary is None:
        return {}
    try:
        result = subprocess.run(
            [binary, "-j", str(path)],
            capture_output=True,
            check=True,
        )
    except (subprocess.CalledProcessError, OSError) as exc:
        logger.warning("exiftool -j failed on %s: %s", path, exc)
        return {}
    try:
        parsed = json.loads(result.stdout.decode("utf-8", "replace"))
    except json.JSONDecodeError as exc:
        logger.warning("could not parse exiftool JSON for %s: %s", path, exc)
        return {}
    if isinstance(parsed, list) and parsed and isinstance(parsed[0], dict):
        return dict(parsed[0])
    return {}


def _read_docx_props(path: Path) -> dict:
    """Read python-docx core properties as a flat lowercase dict."""
    try:
        import docx  # type: ignore[import-not-found]
    except ImportError as exc:
        logger.warning("python-docx not installed; cannot read .docx properties: %s", exc)
        return {}
    cp = docx.Document(str(path)).core_properties
    return {
        "author": cp.author or "",
        "last_modified_by": cp.last_modified_by or "",
        "title": cp.title or "",
        "subject": cp.subject or "",
        "keywords": cp.keywords or "",
        "comments": cp.comments or "",
        "category": cp.category or "",
    }


def _read_xlsx_props(path: Path) -> dict:
    """Read openpyxl workbook properties as a flat dict."""
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ImportError as exc:
        logger.warning("openpyxl not installed; cannot read .xlsx properties: %s", exc)
        return {}
    props = openpyxl.load_workbook(str(path)).properties
    return {
        "creator": props.creator or "",
        "last_modified_by": props.lastModifiedBy or "",
        "title": props.title or "",
        "subject": props.subject or "",
        "keywords": props.keywords or "",
        "description": props.description or "",
        "category": props.category or "",
    }


def read_metadata(path: str | Path) -> dict:
    """Read metadata from a file for verification.

    Uses ``exiftool -j`` when available; for .docx/.xlsx also reads the office
    core properties via the office libraries. Office-derived keys take
    precedence over exiftool's so the lowercase ``author``/``creator`` fields
    reflect the authoritative source.

    Args:
        path: The file to inspect.

    Returns:
        A flat dict of metadata. Empty when nothing could be read.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    meta: dict = _read_exiftool(path)
    if suffix == ".docx":
        meta.update(_read_docx_props(path))
    elif suffix == ".xlsx":
        meta.update(_read_xlsx_props(path))
    return meta
